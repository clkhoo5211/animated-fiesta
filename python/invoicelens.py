"""InvoiceLens — Python / command-line version of the web app (site/index.html).

Same prompt, checks, A/B cross-verification, reconciliation and duplicate detection as the
browser version, for batch / server use. Talks to any OpenAI-compatible endpoint (OpenAI,
Gemini, OpenRouter, Qwen, relays…) or Anthropic over plain HTTPS.

    pip install -r python/requirements.txt
    export IL_A_KEY=sk-...                       # model A key (or --a-key)
    python python/invoicelens.py invoices/*.jpg register.pdf \
        --a-base https://api.openai.com/v1 --a-model gpt-4o --out out/

Outputs in --out: results.json, summary.csv, fields.csv (every value), tables.csv, reconcile.csv
"""
from __future__ import annotations

import argparse, asyncio, base64, csv, hashlib, io, json, os, re, sys, time
from dataclasses import dataclass, field
from pathlib import Path

import httpx
from PIL import Image, ImageFilter, ImageOps

PROMPT = "You are a document data extractor for ANY business document (invoice, receipt, delivery order, invoice register / shipment manifest, purchase order, statement, form...). The page may be rotated or photographed at an angle: read it in its correct orientation.\nReturn ONE JSON object with this shape (null when absent; all values as strings exactly as printed):\n{\"document_type\":\"invoice|receipt|delivery_order|invoice_register|purchase_order|statement|quotation|credit_note|debit_note|packing_list|aging_report|other\",\n \"title\":\"heading as printed\",\"document_number\",\"document_date\",\"currency\",\n \"parties\":[{\"role\":\"supplier|customer|bill_to|ship_to|transporter|issuer|other\",\"name\",\"registration_no\",\"tax_id\",\"address\",\"contact\"}],\n \"fields\":[{\"label\":\"label exactly as printed\",\"value\":\"value as printed\"}],\n \"tables\":[{\"name\",\"columns\":[{\"name\":\"header as printed\",\"role\":\"text|id|date|qty|unit_price|amount|number|row_total|debit|credit|balance\"}],\n   \"rows\":[[\"cell\", \"...\"]],\"row_kinds\":[\"line|group_header|subtotal|other per row, same length as rows\"]|null,\"total_row\":[\"cell or null per column\"]|null,\"printed_row_count\":\"e.g. 7 from '7 Orders'\"|null}],\n \"totals\":[{\"label\",\"value\"}],\"grand_total\",\"amount_in_words\",\n \"stamps_and_chops\":[{\"text\",\"position\"}],\"handwritten_notes\":[{\"text\",\"position\"}],\n \"rotation_degrees\":\"0|90|180|270 = clockwise turn needed to make the text upright\",\n \"low_confidence_fields\":[\"path or label of anything you are unsure about\"]}\nRules:\n- Only output text that is actually on the page. Never invent company names, numbers or rows.\n- Put EVERY labelled value on the page into \"fields\" (one entry per label), even if also used elsewhere. A label with a blank value gets value null.\n- Copy IDs, phone numbers, tax IDs and amounts character by character. Unreadable character -> \"?\" and list the field in low_confidence_fields.\n- Table column headers are not field labels: never repeat table cells in fields.\n- Line items (products/services with qty, price or amount) ALWAYS go in tables, even when the table has no ruled lines; never put per-item amounts in totals. totals is only for summary lines printed once (subtotal, discount, tax, rounding, total payable).\n- Tables: one row per printed row, cells in column order; put the printed totals line in total_row, not in rows. Column role: qty = quantity, unit_price = price per unit, amount = qty x unit_price, row_total = sum of the other numeric cells in that row, debit / credit = money out / in on a statement, balance = running balance, number = other numbers. A column headed Price/Total that already holds the line total (quantity x it would not give the subtotal) is amount, not unit_price.\n- Grouped reports (a group/customer header row, its lines, then a subtotal row such as \"Total for X\"): keep header, line and subtotal rows in rows, in printed order, and mark each in row_kinds. Aging reports: bucket columns (Current, 1-30, 31-60 ... 90+) are number and the outstanding/total column is row_total.\n- Handwriting and tick marks: transcribe literally ONLY in handwritten_notes and list them in low_confidence_fields. If handwriting or a signature cannot be read with confidence, write \"[illegible]\" or \"[signature]\" - never guess names or words. Printed stamp text that overlaps handwriting belongs to stamps_and_chops. Never put handwriting into fields, parties, tables or totals: a printed label whose space is blank, or only has handwriting written over/next to it, gets value null.\n- Letter O vs digit 0, I/l vs 1, S vs 5, B vs 8: decide from context (account numbers, IDs and phone numbers are mostly digits; email domains are real words such as jaring.my, gmail.com).\n- A party's name is the name only: a registration or tax number printed next to it goes in registration_no / tax_id, not in name.\n- registration_no is only a company registration number printed as such (e.g. \"(279018-W)\", \"Co. No.\", \"Reg. No.\", \"SSM\"). Branch, store, outlet or customer codes are not registration numbers: put them in fields.\n- Parties come only from printed letterheads / address blocks (e.g. \"Billing Address\", \"Delivery Address\", \"Bill To\"). A rubber stamp, company chop or \"Received by\" stamp is NOT a party: put it only in stamps_and_chops.\n- Labels: read each small label carefully and copy it exactly; never rename it. Digits that belong to a label (e.g. the \"1\" in \"Ref 1:\") are not its value. If the space after a label is empty, its value is null.\nRaw JSON only, no markdown."

MULTI_NOTE = ("\n\nYou receive {n} images of the SAME page: image 1 is the full page, the others are enlarged overlapping "
              "sections (top/bottom or left/right halves) for reading small text. Use the close-ups to read characters; "
              "output each field and table row only ONCE.")
MODEL_TIMEOUT = 180.0
MAX_TEXT = 60000
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}


# ----------------------------------------------------------------- helpers
def num(v):
    """Amounts incl. accounting negatives: '(1,234.56)', '1,234.56-', '1,234.56 CR', '-RM5'."""
    s = str(v if v is not None else "").strip()
    try:
        n = float(re.sub(r"[^0-9.]", "", s))
    except ValueError:
        return None
    neg = bool(re.match(r"^[^0-9]*-", s) or re.fullmatch(r"\s*\(.*\)\s*", s) or re.search(r"\d\s*-\s*$", s) or re.search(r"\bCR\.?\s*$", s, re.I))
    return -n if neg else n


def norm_text(s):
    """Invisible / look-alike characters from PDF text layers and spreadsheets that break IDs."""
    s = re.sub(r"[\u00AD\u200B-\u200D\u2060\uFEFF]", "", str(s or "")).replace("\x00", " ")
    s = re.sub(r"[\u2010-\u2015\u2212\uFE58\uFE63\uFF0D]", "-", s)
    s = re.sub(r"[\u00A0\u2007\u202F\u3000]", " ", s)
    return re.sub(r"[\u201C\u201D\u201F]", '"', re.sub(r"[\u2018\u2019\u201B]", "'", s))


def parse_json(text: str):
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", (text or "").strip())
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", text, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except json.JSONDecodeError:
                pass
        return {"raw": text}


def jpeg_b64(img: Image.Image, quality=95) -> str:
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=quality)
    return base64.b64encode(buf.getvalue()).decode()


def b64_image(b64: str) -> Image.Image:
    return Image.open(io.BytesIO(base64.b64decode(b64)))


# ----------------------------------------------------------------- ingestion
@dataclass
class Segment:
    id: str
    b64: str | None = None          # JPEG page/frame image
    text: str = ""                  # PDF text layer or extracted document text
    kind: str = "image"             # image | text
    source: str = ""
    truncated: bool = False
    raw_rows: list | None = None    # spreadsheet rows (for local parse without a model)
    text_only: bool = False         # PDF page sent as its text layer (image kept for local QR)
    straightened: int = 0           # degrees the page was auto-rotated after the model reported it sideways


def kind_of(path: Path) -> str | None:
    e = path.suffix.lower()
    if e in (".heic", ".heif"): return "heic"
    if e == ".pdf": return "pdf"
    if e in (".xlsx", ".xlsm", ".xls", ".ods"): return "sheet"
    if e == ".csv": return "csv"
    if e == ".docx": return "docx"
    if e in (".mp4", ".mov", ".avi", ".mkv", ".webm"): return "video"
    if e in IMAGE_EXT: return "image"
    return None


def _rows_text(rows):
    return "\n".join(" | ".join(re.sub(r"\s+", " ", str(c)).strip() for c in r) for r in rows)


def ingest(path: Path, rotate: int = 0, pdf_text: str = "auto") -> list[Segment]:
    k = kind_of(path)
    segs: list[Segment] = []
    if k in ("image", "heic"):
        if k == "heic":
            import pillow_heif
            pillow_heif.register_heif_opener()
        img = ImageOps.exif_transpose(Image.open(path))
        segs.append(Segment("Image_1", jpeg_b64(img)))
    elif k == "pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            for i, page in enumerate(doc, 1):
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                text = page.get_text()
                if pdf_text != "image" and len(re.sub(r"\s", "", text)) >= 200:   # real text layer: send text, keep image for QR
                    segs.append(Segment(f"Page_{i}", jpeg_b64(img), text[:MAX_TEXT], "text", "PDF text layer", len(text) > MAX_TEXT, None, True))
                else:
                    segs.append(Segment(f"Page_{i}", jpeg_b64(img), text=text))
    elif k == "video":
        import cv2
        cap = cv2.VideoCapture(str(path))
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        for n, frac in enumerate((0.25, 0.5, 0.75), 1):
            cap.set(cv2.CAP_PROP_POS_FRAMES, int(total * frac))
            ok, frame = cap.read()
            if ok:
                img = Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
                segs.append(Segment(f"Video_Frame_{n}_{int(frac*100)}pct", jpeg_b64(img)))
        cap.release()
    elif k in ("sheet", "csv"):
        sheets = []
        if k == "csv":
            with open(path, newline="", encoding="utf-8-sig") as f:
                sheets.append(("CSV", "CSV", [r for r in csv.reader(f)]))
        else:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            for ws in wb.worksheets:
                rows = [["" if c is None else (f"{c:g}" if isinstance(c, float) else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
                sheets.append((f"Sheet_{ws.title}", f"Excel sheet “{ws.title}”", rows))
        for sid, src, rows in sheets:
            rows = [r for r in rows if any(str(c).strip() for c in r)]
            if not rows:
                continue
            txt = _rows_text(rows)
            segs.append(Segment(sid, None, txt[:MAX_TEXT], "text", src, len(txt) > MAX_TEXT, rows))
    elif k == "docx":
        import docx
        d = docx.Document(path)
        parts = []
        body = d.element.body
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                t = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t")).strip()
                if t:
                    parts.append(t)
            elif tag == "tbl":
                tbl = docx.table.Table(child, d)
                parts.append("[TABLE]\n" + _rows_text([[c.text for c in r.cells] for r in tbl.rows]) + "\n[/TABLE]")
        text = "\n".join(parts).strip()
        if not text:
            raise ValueError("The Word document has no readable text (it may only contain images)")
        segs.append(Segment("Document", None, text[:MAX_TEXT], "text", "Word document", len(text) > MAX_TEXT))
    else:
        raise ValueError(f"Unsupported file type: {path.name}")
    if rotate:
        for s in segs:
            if s.b64:
                s.b64 = jpeg_b64(b64_image(s.b64).rotate(-rotate, expand=True))
    return segs


# ----------------------------------------------------------------- image prep / local barcodes
def enhance(img: Image.Image, box=None, target=2000, do_enhance=True) -> str:
    if box:
        img = img.crop(box)
    k = min(3.0, max(1.0, target / max(img.size)))
    img = img.resize((round(img.width * k), round(img.height * k)), Image.LANCZOS)
    if do_enhance:
        img = ImageOps.autocontrast(ImageOps.grayscale(img), cutoff=1).filter(ImageFilter.UnsharpMask(radius=2, percent=160, threshold=2))
    return jpeg_b64(img)


def prep_images(b64: str, do_enhance: bool, tiles: bool) -> list[str]:
    if not do_enhance and not tiles:
        return [b64]
    img = b64_image(b64)
    out = [enhance(img, None, 2000, do_enhance)]
    if tiles:
        W, H = img.size
        ov = 0.08
        if H >= W:
            h = round(H * (0.5 + ov))
            out += [enhance(img, (0, 0, W, h), 2000, do_enhance), enhance(img, (0, H - h, W, H), 2000, do_enhance)]
        else:
            w = round(W * (0.5 + ov))
            out += [enhance(img, (0, 0, w, H), 2000, do_enhance), enhance(img, (W - w, 0, W, H), 2000, do_enhance)]
    return out


QR_REGIONS = [(0, .5, .55, .5), (.45, .5, .55, .5), (0, 0, .55, .5), (.45, 0, .55, .5), (0, .25, .55, .5), (.45, .25, .55, .5)]


def local_codes(seg: Segment) -> dict:
    if not seg.b64:
        return {"qr_and_barcodes": [], "native_extracted_text": f"({seg.source}, {len(seg.text)} chars{', truncated' if seg.truncated else ''})", "not_applicable": True}
    try:
        import zxingcpp
        img = b64_image(seg.b64)
        found = zxingcpp.read_barcodes(img)
        for scale in (2, 3):
            if found:
                break
            big = ImageOps.autocontrast(ImageOps.grayscale(img)).resize((img.width * scale, img.height * scale), Image.NEAREST)
            found = zxingcpp.read_barcodes(big)
        # small codes in a corner of a low-resolution photo: scan overlapping regions, enlarged
        for rx, ry, rw, rh in QR_REGIONS:
            if found:
                break
            box = (int(img.width * rx), int(img.height * ry), int(img.width * (rx + rw)), int(img.height * (ry + rh)))
            crop = ImageOps.autocontrast(ImageOps.grayscale(img.crop(box)))
            for k, rs, us in ((3, Image.BICUBIC, (2, 300, 1)), (4, Image.LANCZOS, (4, 200, 2))):  # smooth enlarge + sharpen
                found = found or zxingcpp.read_barcodes(crop.resize((crop.width * k, crop.height * k), rs).filter(ImageFilter.UnsharpMask(*us)))
        return {"qr_and_barcodes": [{"type": str(r.format).split(".")[-1], "data": r.text} for r in found], "native_extracted_text": seg.text}
    except Exception as e:  # pragma: no cover - local decoding is best effort
        return {"qr_and_barcodes": [], "native_extracted_text": seg.text, "error": str(e)}


# ----------------------------------------------------------------- model calls
@dataclass
class ModelConfig:
    type: str = "openai"            # openai (any OpenAI-compatible) | anthropic
    base: str = "https://api.openai.com/v1"
    model: str = "gpt-4o"
    key: str = ""
    json_mode: bool = True
    timeout: float = MODEL_TIMEOUT
    max_tokens: int = 32000

    @property
    def enabled(self):
        return bool(self.key)


def clean_base(u: str) -> str:
    u = (u or "").strip().rstrip("/")
    return re.sub(r"/(models|chat/completions)$", "", u)


def _with_usage(r, i, o):
    if isinstance(r, dict) and (i or o):
        r["__usage"] = {"in": i or 0, "out": o or 0}
    return r


RETRY_CODES = {429, 502, 503, 504}
TEMP_MSG = re.compile(r"capacity|temporarily unavailable|overloaded|try again|retry shortly|rate.?limit", re.I)
OUT_OF_TOKENS = "Model used all {n} output tokens (probably on thinking) and returned no answer - raise --max-tokens"
RETRY_DELAYS = [5.0, 15.0, 30.0]
RETRY_BUDGET = 300.0  # stop retrying once a call has taken 5 min in total


class HTTPStatusError(RuntimeError):
    def __init__(self, msg, status):
        super().__init__(msg)
        self.status = status


async def call_model(client: httpx.AsyncClient, c: ModelConfig, text: str, images: list[str] | None, max_tokens=None):
    """Call the model, retrying temporary errors (429/502/503/504) with backoff."""
    t0 = time.monotonic()
    for i in range(len(RETRY_DELAYS) + 1):
        try:
            return await _call_model_once(client, c, text, images, max_tokens or c.max_tokens)
        except HTTPStatusError as e:
            if not (e.status in RETRY_CODES or TEMP_MSG.search(str(e))) or i >= len(RETRY_DELAYS) \
                    or time.monotonic() - t0 + RETRY_DELAYS[i] > RETRY_BUDGET:
                raise
            print(f"  temporary error {e.status} from {c.model}, retrying ({i + 1}/{len(RETRY_DELAYS)})", file=sys.stderr)
            await asyncio.sleep(RETRY_DELAYS[i])


async def _call_model_once(client: httpx.AsyncClient, c: ModelConfig, text: str, images: list[str] | None, max_tokens: int):
    images = images or []
    if c.type == "anthropic":
        content = [{"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": d}} for d in images] + [{"type": "text", "text": text}]
        r = await client.post("https://api.anthropic.com/v1/messages", timeout=c.timeout,
                              headers={"x-api-key": c.key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                              json={"model": c.model, "max_tokens": max_tokens, "messages": [{"role": "user", "content": content}]})
        try:
            j = r.json() if r.content else {}
        except ValueError:
            j = {}
        if r.status_code >= 400:
            raise HTTPStatusError((j.get("error") or {}).get("message") or f"HTTP {r.status_code}", r.status_code)
        u = j.get("usage") or {}
        txt = next((b["text"] for b in j.get("content", []) if b.get("type") == "text"), "")
        if not txt.strip() and j.get("stop_reason") == "max_tokens":
            raise RuntimeError(OUT_OF_TOKENS.format(n=max_tokens))
        return _with_usage(parse_json(txt), u.get("input_tokens"), u.get("output_tokens"))
    content = [{"type": "text", "text": text}] + [{"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{d}", "detail": "high"}} for d in images]
    body = {"model": c.model, "max_tokens": max_tokens, "messages": [{"role": "user", "content": content}]}
    if c.json_mode:
        body["response_format"] = {"type": "json_object"}
    r = await client.post(clean_base(c.base) + "/chat/completions", timeout=c.timeout,
                          headers={"Authorization": f"Bearer {c.key}", "content-type": "application/json"}, json=body)
    try:
        j = r.json() if r.content else {}
    except ValueError:
        j = {}
    if r.status_code >= 400:
        raise HTTPStatusError((j.get("error") or {}).get("message") or f"HTTP {r.status_code}", r.status_code)
    u = j.get("usage") or {}
    ch = (j.get("choices") or [{}])[0]
    txt = (ch.get("message") or {}).get("content") or ""
    if not txt.strip() and ch.get("finish_reason") == "length":
        raise RuntimeError(OUT_OF_TOKENS.format(n=max_tokens))
    return _with_usage(parse_json(txt), u.get("prompt_tokens"), u.get("completion_tokens"))


async def safe(coro):
    try:
        return await coro
    except httpx.TimeoutException:
        return {"error": "Model did not reply in time (timeout)"}
    except Exception as e:
        return {"error": str(e)}


# ----------------------------------------------------------------- checks (same rules as the web app)
# grouped reports: the model marks rows in row_kinds; the label pattern is only a fallback
SUBROW = re.compile(r"^\s*(sub.?total|total\s+for\b|total\s*[-–:]|customer\s+total|group\s+total|jumlah\s+kecil|小计)", re.I)


def group_sums(rows, numeric, cols, ti, tn, add, is_sub):
    """Each subtotal row must equal the sum of the lines since the previous subtotal."""
    if not any(is_sub):
        return
    acc, start, bad, n = [0.0] * len(numeric), 0, 0, 0
    for ri, row in enumerate(rows):
        if is_sub[ri]:
            n += 1
            for k, ci in enumerate(numeric):
                v = num(row[ci]) if ci < len(row) else None
                if v is not None and abs(v - acc[k]) > 0.011:
                    bad += 1
                    add(False, "groupsum", ti=ti, tn=tn, row=ri, col=cols[ci].get("name"), sum=round(acc[k], 2), printed=row[ci], **{"from": start + 1})
            acc, start = [0.0] * len(numeric), ri + 1
        else:
            for k, ci in enumerate(numeric):
                acc[k] += (num(row[ci]) or 0) if ci < len(row) else 0
    if not bad:
        add(True, "groupsum_all", ti=ti, tn=tn, n=n)


def grounding(layer, r):
    """Values the model read that should be printed verbatim (IDs, dates, amounts) must appear in the PDF text layer."""
    if not layer or not isinstance(r, dict):
        return None
    vals = []
    def put(path, v):
        v = str(v if v is not None else "").strip()
        if re.search(r"\d", v) and len(re.sub(r"\D", "", v)) >= 3:
            vals.append((path, v))
    for k in ("document_number", "document_date", "grand_total"):
        put(k, r.get(k))
    for i, p in enumerate(x for x in r.get("parties") or [] if isinstance(x, dict)):
        put(f"parties.{i}.registration_no", p.get("registration_no")); put(f"parties.{i}.tax_id", p.get("tax_id"))
    for key in ("fields", "totals"):
        for i, f in enumerate(r.get(key) or []):
            put(f"{key}.{i}.value", (f or {}).get("value"))
    flat = re.sub(r"\s+", "", layer).upper()
    nums = {num(x) for x in re.findall(r"\d[\d,]*\.\d+", layer)}
    miss = [{"field": p, "value": v} for p, v in vals
            if not (re.sub(r"\s+", "", v).upper() in flat or (re.search(r"\d\.\d{2}\b", v) and num(v) is not None and abs(num(v)) in nums))]
    return {"checked": len(vals), "missing": miss}


def verify(r) -> dict | None:
    if not isinstance(r, dict) or r.get("error") or r.get("status"):
        return None
    checks = []
    add = lambda ok, code, **p: checks.append({"ok": bool(ok), "code": code, **p})
    for ti, t in enumerate(r.get("tables") or []):
        cols = t.get("columns") or []
        rows = [x for x in (t.get("rows") or []) if isinstance(x, list)]
        tn = t.get("name")
        idx = lambda role: [i for i, c in enumerate(cols) if (c or {}).get("role") == role]
        numeric = [i for i, c in enumerate(cols) if (c or {}).get("role") in ("qty", "amount", "number", "row_total")]
        kinds = t.get("row_kinds") if isinstance(t.get("row_kinds"), list) else None
        is_sub = [(kinds[i] == "subtotal") if kinds and i < len(kinds) else (not kinds and any(isinstance(c, str) and SUBROW.search(c) for c in row))
                  for i, row in enumerate(rows)]
        group_sums(rows, numeric, cols, ti, tn, add, is_sub)
        tot = t.get("total_row")
        if isinstance(tot, list):
            for ci in numeric:
                tv = num(tot[ci]) if ci < len(tot) else None
                if tv is None:
                    continue
                s = sum(num(row[ci]) or 0 for i, row in enumerate(rows) if ci < len(row) and not is_sub[i])
                add(abs(s - tv) < 0.011, "colsum", ti=ti, tn=tn, col=cols[ci].get("name"), sum=round(s, 2), printed=tot[ci])
        rt = idx("row_total")
        parts = [i for i in numeric if i not in rt and cols[i].get("role") != "amount"]
        bad_rows = 0
        if len(rt) == 1 and parts:
            for ri, row in enumerate(rows):
                tv = num(row[rt[0]]) if rt[0] < len(row) else None
                if tv is None:
                    continue
                s = sum(num(row[i]) or 0 for i in parts if i < len(row))
                if abs(s - tv) > 0.011:
                    bad_rows += 1
                    add(False, "rowsum", ti=ti, tn=tn, row=ri, expr=f"{'+'.join(cols[i].get('name','') for i in parts)} = {round(s,2):g} ≠ {row[rt[0]]}")
            if not bad_rows:
                add(True, "rowsum_all", ti=ti, tn=tn, n=len(rows))
        q, pz, am = (idx("qty") or [None])[0], (idx("unit_price") or [None])[0], (idx("amount") or [None])[0]
        if None not in (q, pz, am):
            bad = 0
            for ri, row in enumerate(rows):
                a, b, c = (num(row[i]) if i < len(row) else None for i in (q, pz, am))
                if None in (a, b, c):
                    continue
                if abs(abs(a * b) - abs(c)) > 0.011:  # credit notes print negative amounts
                    bad += 1
                    add(False, "amount", ti=ti, tn=tn, row=ri, expr=f"{row[q]} × {row[pz]} = {a*b:.2f} ≠ {row[am]}")
            if not bad:
                add(True, "qtyprice_all", ti=ti, tn=tn, n=len(rows))
        if len(rows) >= 3:
            last = rows[-1]
            big = [ci for ci in numeric if ci < len(last) and (num(last[ci]) or 0) > 0
                   and num(last[ci]) > sum(num(r[ci]) or 0 for r in rows[:-1] if ci < len(r))]
            tr = t.get("total_row") if isinstance(t.get("total_row"), list) else []
            if len(big) >= 2 and not any(num(tr[ci]) for ci in numeric if ci < len(tr)):
                add(False, "lastrow", ti=ti, tn=tn, row=len(rows) - 1, cols=", ".join(cols[i].get("name") or "" for i in big))
        pc = num(t.get("printed_row_count"))
        if pc is not None:
            add(pc == len(rows), "rowcount", ti=ti, tn=tn, printed=t.get("printed_row_count"), n=len(rows))
    gt = num(r.get("grand_total"))
    if gt is not None:
        amts = [sum(num(row[i]) or 0 for row in (t.get("rows") or []) if isinstance(row, list) and i < len(row))
                for t in (r.get("tables") or []) for i, c in enumerate(t.get("columns") or []) if (c or {}).get("role") == "amount"]
        if len(amts) == 1:
            tot = r.get("totals") or []
            tax = sum(num(x.get("value")) or 0 for x in tot if re.search(r"tax|gst|sst|vat|cukai|税", x.get("label") or "", re.I))
            disc = sum(num(x.get("value")) or 0 for x in tot if re.search(r"disc|diskaun|折", x.get("label") or "", re.I))
            add(abs(amts[0] + tax - disc - gt) < 0.011, "grand", lines=f"{amts[0]:.2f}", tax=tax, disc=disc, total=r.get("grand_total"))
    totals_checks(r, add)
    date_order(r, add)
    tax_reg_check(r, add)
    guessed_names(r, add)
    phone_lengths(r, add)
    for ti, t in enumerate(r.get("tables") or []):
        balance_check(t, ti, add)
    cross_checks(r, add)
    items_as_totals(r, add)
    return {"passed": all(c["ok"] or c.get("warn") for c in checks), "checks": checks}


def _int(v):
    try:
        return int(str(v).strip().split()[0])
    except (ValueError, IndexError):
        return None


def _days(a, b):
    import datetime
    try:
        return (datetime.date.fromisoformat(b) - datetime.date.fromisoformat(a)).days
    except ValueError:
        return None


def _texts(items):
    return [x if isinstance(x, str) else (x or {}).get("text") or "" for x in items or []]


def cross_checks(r, add):
    """Advisory checks: handwritten/stamp dates earlier than the document date; stamp number one digit off an address number."""
    doc = norm_date(r.get("document_date"))
    if doc:
        for n in _texts((r.get("handwritten_notes") or []) + (r.get("stamps_and_chops") or [])):
            for m in re.findall(r"\b\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}\b", n):
                d = norm_date(m)
                if d and d < doc and _days(d, doc) is not None and _days(d, doc) <= 366:
                    add(False, "notedate", warn=True, v=m, doc=r.get("document_date"), note=n[:60])
    addr = " ".join(f"{p.get('name') or ''} {p.get('address') or ''}" for p in r.get("parties") or [] if isinstance(p, dict))
    an = set(re.findall(r"\d{3,6}", addr))
    for st in _texts(r.get("stamps_and_chops")):
        for x in set(re.findall(r"\b\d{3,6}\b", st)):
            if x in an:
                continue
            y = next((a for a in sorted(an) if len(a) == len(x) and sum(p != q for p, q in zip(a, x)) == 1), None)
            if y:
                add(False, "stampno", warn=True, v=x, a=y)
    hw = [re.sub(r"\s+", " ", n).strip().lower() for n in _texts(r.get("handwritten_notes"))]
    hw = [h for h in hw if len(h) >= 4]
    for f in r.get("fields") or []:
        v = re.sub(r"\s+", " ", str((f or {}).get("value") or "")).strip().lower()
        if len(v) >= 4 and any(h == v or v in h or h in v for h in hw):
            add(False, "hwfield", warn=True, l=f.get("label"), v=f.get("value"))
    for p in r.get("parties") or []:
        m = re.match(r"^(?:TIN:?)?C(\d+)$", re.sub(r"\s", "", str((p or {}).get("tax_id") or "")), re.I)
        if m and len(m.group(1)) not in (10, 11):
            add(False, "tin", warn=True, v=p.get("tax_id"), n=len(m.group(1)), who=p.get("name") or p.get("role"))
    blob = json.dumps([r.get("parties"), r.get("fields")], ensure_ascii=False)
    for em in sorted(set(re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", blob))):
        d = re.sub(r"^(po|mail)\.", "", em.split("@")[1].lower())
        near = None if d in KNOWN_DOMAINS else next((k for k in KNOWN_DOMAINS if _edit1(d, k)), None)
        if near:
            add(False, "email", warn=True, v=em, k=near)
    for f in r.get("fields") or []:
        v = str((f or {}).get("value") or "")
        if re.search(r"(no|number|account|acc|code|kod|akaun|ref)\b", f.get("label") or "", re.I) and re.fullmatch(r"[A-Z0-9-]{3,}", v) \
                and re.search(r"\d", v) and re.search(r"(?<=[A-Z0-9])O(?=\d)|(?<=\d)O", v):
            add(False, "o0", warn=True, l=f.get("label"), v=v, s=re.sub(r"(?<=[A-Z0-9])O(?=\d)|(?<=\d)O", "0", v))


def items_as_totals(r, add):
    """Line items that ended up in totals: a repeated label, or 3+ totals summing to the grand total with no item table."""
    tot = [x for x in r.get("totals") or [] if isinstance(x, dict) and num(x.get("value")) is not None]
    if len(tot) < 2:
        return
    labs = {}
    for x in tot:
        k = str(x.get("label") or "").strip().lower()
        labs[k] = labs.get(k, 0) + 1
    rep = next(((k, n) for k, n in labs.items() if k and n >= 2), None)
    has_amt = any(any((c or {}).get("role") == "amount" for c in t.get("columns") or []) and t.get("rows") for t in r.get("tables") or [])
    gt = num(r.get("grand_total"))
    cand = [x for x in tot if (gt is None or abs(num(x["value"]) - gt) > 0.011) and not SUMMARY_LBL.search(x.get("label") or "")]
    s = sum(num(x["value"]) for x in cand)
    if rep or (not has_amt and gt is not None and len(cand) >= 3 and abs(s - gt) < 0.011):
        add(False, "itemstot", l=rep[0] if rep else "", n=rep[1] if rep else len(cand), sum=round(s, 2))


def date_order(r, add):
    """03/04/2026 is 3 April or 4 March: other dates on the page settle it (a part > 12); otherwise warn."""
    vals = [r.get("document_date")] + [(f or {}).get("value") for f in r.get("fields") or []]
    for t in r.get("tables") or []:
        di = [i for i, c in enumerate(t.get("columns") or []) if (c or {}).get("role") == "date"]
        vals += [row[i] for row in t.get("rows") or [] if isinstance(row, list) for i in di if i < len(row)]
    rx = r"\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})\b"
    dmy = mdy = 0
    for v in vals:
        for a, b, _ in re.findall(rx, str(v or "")):
            a, b = int(a), int(b)
            if a > 12 >= b: dmy += 1
            elif b > 12 >= a: mdy += 1
    dd = re.search(rx, str(r.get("document_date") or ""))
    if mdy and not dmy:
        add(False, "datemdy", warn=True, n=mdy)
    elif dd and not dmy and re.search(r"\bUSD\b|US\$|^\s*\$", f"{r.get('currency') or ''} {r.get('grand_total') or ''}", re.I) and int(dd.group(1)) <= 12 and int(dd.group(2)) <= 12 and dd.group(1) != dd.group(2):
        a, b = int(dd.group(1)), int(dd.group(2))
        add(False, "dateamb", warn=True, v=r.get("document_date"), a=f"{a}/{b}", b=f"{b}/{a}")


def _lbl(l):
    return re.sub(r"[\W_]+", "", str(l or "").lower())


def align_by_label(a, b):
    """Put B's fields / totals in A's order (matched by label) so A/B comparison shows real differences, not ordering."""
    if not isinstance(a, dict) or not isinstance(b, dict) or a.get("error") or b.get("error"):
        return
    for k in ("fields", "totals"):
        A, B = a.get(k), b.get(k)
        if not isinstance(A, list) or not isinstance(B, list):
            continue
        pool, out = list(B), []
        for x in A:
            i = next((j for j, y in enumerate(pool) if isinstance(y, dict) and isinstance(x, dict) and _lbl(y.get("label")) == _lbl(x.get("label"))), -1)
            out.append(pool.pop(i) if i >= 0 else None)
        while out and out[-1] is None:
            out.pop()
        b[k] = [x if x is not None else {"label": None, "value": None} for x in out] + pool


NAME_RX = re.compile(r"^[A-Z][A-Za-z'’.-]+(\s+(bin|binti|bt|b\.|a/l|a/p|al|ap|van|de|[A-Z][A-Za-z'’.-]+)){1,5}$")


def guessed_names(r, add):
    """A full personal name in handwriting that the model itself was unsure about is most likely invented."""
    low = " ".join(str(x) for x in r.get("low_confidence_fields") or [])
    for i, x in enumerate(r.get("handwritten_notes") or []):
        v = str(x.get("text") if isinstance(x, dict) else x or "").strip()
        if NAME_RX.match(v) and not v.startswith("[") and (re.search(rf"handwritten_notes\W*{i}\b", low) or v in low):
            add(False, "guessname", v=v)


def phone_lengths(r, add):
    """Phone / fax numbers with the same area code should have the same number of digits."""
    txt = json.dumps([r.get("parties"), r.get("fields")], ensure_ascii=False)
    by = {}
    for m in re.finditer(r"(?<![\d-])(\+?\d{1,4})[-\s](\d{3,5})[\s-]?(\d{3,5})(?![\d-])", txt):
        n = len(m.group(2) + m.group(3))
        if 6 <= n <= 10:
            by.setdefault(m.group(1), []).append((m.group(0), n))
    for pre, lst in by.items():
        cnt = lambda n: sum(1 for _, k in lst if k == n)
        ns = sorted({n for _, n in lst}, key=cnt, reverse=True)
        if len(ns) < 2:
            continue
        if cnt(ns[0]) > cnt(ns[1]):  # a clear majority length: flag the odd ones
            for v, n in lst:
                if n != ns[0]:
                    add(False, "phonelen", warn=True, v=v, p=pre, n=n, c=ns[0])
        else:  # no majority: show them side by side
            add(False, "phonemix", warn=True, p=pre, v=" · ".join(f"{v} ({n})" for v, n in lst))


def tax_reg_check(r, add):
    """Tax charged but no tax registration printed; a registration number repeated inside a party name."""
    tax = [x for x in r.get("totals") or [] if isinstance(x, dict) and re.search(r"tax|sst|gst|vat|cukai|税", x.get("label") or "", re.I) and (num(x.get("value")) or 0) > 0]
    parties = [p for p in r.get("parties") or [] if isinstance(p, dict)]
    sellers = [p for p in parties if re.search(r"supplier|issuer", p.get("role") or "")]
    has_no = any(p.get("tax_id") for p in (sellers or parties))
    field_no = any(re.search(r"(sst|gst|vat|tax|tin|cukai|税)", f.get("label") or "", re.I) and re.search(r"(reg|no\b|no\.|number|id\b|tin|登记|号)", f.get("label") or "", re.I)
                   and re.search(r"\d", str(f.get("value") or "")) for f in r.get("fields") or [] if isinstance(f, dict))
    if tax and not has_no and not field_no:
        add(False, "taxnoreg", warn=True, l=tax[0].get("label"), v=tax[0].get("value"))
    for p in parties:
        if p.get("name") and p.get("registration_no") and str(p["registration_no"]) in str(p["name"]):
            add(False, "namereg", warn=True, n=p["name"], r=p["registration_no"])


def totals_checks(r, add):
    """Totals block: subtotal + tax - discount + rounding = grand total; 'SST 8%' = 8% of subtotal."""
    tot = [x for x in r.get("totals") or [] if isinstance(x, dict) and x.get("label") and num(x.get("value")) is not None]
    f = lambda rx: [x for x in tot if re.search(rx, x["label"], re.I)]
    sub, gt = next(iter(f(r"sub.?total|jumlah kecil|小计")), None), num(r.get("grand_total"))
    if not sub or gt is None:
        return
    s = num(sub["value"])
    tax, disc, rnd = f(r"tax|sst|gst|vat|cukai|税"), f(r"disc|diskaun|折"), f(r"round|pembundaran|舍入")
    calc = s + sum(num(x["value"]) for x in tax) - sum(abs(num(x["value"])) for x in disc) + sum(num(x["value"]) for x in rnd)
    add(abs(calc - gt) < 0.011, "totalsum", s=sub["value"], c=round(calc, 2), g=r.get("grand_total"))
    for x in tax:
        m = re.search(r"(\d+(?:\.\d+)?)\s*%", x["label"])
        if m:
            want = s * float(m.group(1)) / 100
            add(abs(want - num(x["value"])) < 0.051, "taxrate", l=x["label"], v=x["value"], w=f"{want:.2f}")


def balance_check(t, ti, add):
    """Statements: previous balance -/+ debit/credit = this balance (either sign convention)."""
    cols = t.get("columns") or []
    rows = [r for r in t.get("rows") or [] if isinstance(r, list)]
    ix = lambda role: next((i for i, c in enumerate(cols) if (c or {}).get("role") == role), -1)
    b, d, c = ix("balance"), ix("debit"), ix("credit")
    if b < 0 or (d < 0 and c < 0) or len(rows) < 2:
        return
    g = lambda row, i: num(row[i]) if 0 <= i < len(row) else None
    bal = [g(r, b) for r in rows]
    mv = [(g(r, c) or 0) - (g(r, d) or 0) for r in rows]
    bad, sign = [], None
    for i in range(1, len(rows)):
        if bal[i] is None or bal[i - 1] is None:
            continue
        delta = round(bal[i] - bal[i - 1], 2)
        if sign is None:
            if abs(delta - mv[i]) < 0.011: sign = 1
            elif abs(delta + mv[i]) < 0.011: sign = -1
            else:
                bad.append(i); continue
        if abs(delta - sign * mv[i]) > 0.011:
            bad.append(i)
    for i in bad:
        add(False, "balance", ti=ti, tn=t.get("name"), row=i, p=rows[i - 1][b], v=rows[i][b])
    if not bad:
        add(True, "balance_all", ti=ti, tn=t.get("name"), n=len(rows))


SUMMARY_LBL = re.compile(r"sub|tax|sst|gst|vat|disc|round|cukai|diskaun|小计|税|折|總|总计", re.I)


KNOWN_DOMAINS = ["gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "live.com", "icloud.com", "jaring.my", "streamyx.com", "tm.net.my", "yahoo.com.my", "gmail.com.my"]


def _edit1(a, b):
    """True when a and b differ by exactly one insert, delete or substitution."""
    if a == b or abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) == 1
    s, l = sorted((a, b), key=len)
    return any(l[:i] + l[i + 1:] == s for i in range(len(l)))


def _flatten(o, p="", out=None):
    out = {} if out is None else out
    if isinstance(o, dict):
        for k, v in o.items():
            if k != "low_confidence_fields":
                _flatten(v, f"{p}.{k}" if p else k, out)
    elif isinstance(o, list):
        for i, v in enumerate(o):
            _flatten(v, f"{p}.{i}" if p else str(i), out)
    elif o is not None and o != "":
        out[p] = o
    return out


def compare(a, b):
    if not isinstance(a, dict) or not isinstance(b, dict) or any(x.get("error") or x.get("status") for x in (a, b)):
        return None
    norm = lambda v: re.sub(r"[\s,.:;()\-_/\"']", "", str(v if v is not None else "").lower())
    A, B = _flatten(a), _flatten(b)
    keys = [k for k in dict.fromkeys([*A, *B]) if not re.match(r"^(stamps_and_chops|handwritten_notes)", k)]
    agree, diff = [], []
    for k in keys:
        if k not in A or k not in B:
            diff.append({"field": k, "a": A.get(k), "b": B.get(k), "kind": "missing"}); continue
        nx, ny = num(A[k]), num(B[k])
        if norm(A[k]) == norm(B[k]) or (nx is not None and ny is not None and re.fullmatch(r"[\d.,\s-]+", str(A[k])) and abs(nx - ny) < 0.005):
            agree.append(k)
        else:
            diff.append({"field": k, "a": A[k], "b": B[k], "kind": "mismatch"})
    return {"agreement_rate": round(len(agree) / len(keys), 3) if keys else None, "agreed_fields": len(agree), "discrepancies": diff}


def completeness(r):
    """How much table content a result captured (a truncated description loses a tie)."""
    return sum(len(re.sub(r"\s", "", str(c if c is not None else ""))) for t in r.get("tables") or [] for row in t.get("rows") or [] if isinstance(row, list) for c in row)


def check_score(c):
    """Passed checks count, failures cost more; a result that ran no checks does not win by default."""
    return sum(1 if x["ok"] else (-0.5 if x.get("warn") else -3) for x in c["checks"]) if c else float("-inf")


MERGE_SKIP = re.compile(r"(^|\.)(role|label|row_kinds|columns|name|low_confidence_fields|rotation_degrees|document_type)(\.|$)")


def _get(o, path):
    for k in path.split(".") if path else []:
        if isinstance(o, list) and k.isdigit() and int(k) < len(o): o = o[int(k)]
        elif isinstance(o, dict): o = o.get(k)
        else: return None
    return o


def merge_from(final, who, diffs):
    """Keep the chosen model's result, but take the other model's value where the chosen one is empty or a truncated version of it."""
    if not isinstance(final, dict) or not diffs:
        return []
    other, mine = ("b", "a") if who == "A" else ("a", "b")
    out = []
    s = lambda v: re.sub(r"\s+", " ", str(v if v is not None else "")).strip()
    for d in diffs:
        f, o, m = d["field"], d.get(other), d.get(mine)
        if MERGE_SKIP.search(f) or re.fullmatch(r"tables\.\d+\.name", f) or o is None or isinstance(o, (dict, list)) or s(o) == "":
            continue
        trunc = len(s(o)) > len(s(m)) and s(o).lower().startswith(re.sub(r"[\W_]+$", "", s(m).lower())) and num(m) is None
        parent = _get(final, ".".join(f.split(".")[:-1]))
        if (s(m) == "" or trunc) and isinstance(parent, (dict, list)):
            k = f.split(".")[-1]
            if isinstance(parent, list):
                if not k.isdigit() or int(k) >= len(parent): continue
                parent[int(k)] = o
            else:
                parent[k] = o
            out.append({"path": f, "from": other.upper(), "v": o})
    return out


def pick_model(seg_result):
    a, b = seg_result["provider_responses"]["model_a"], seg_result["provider_responses"]["model_b"]
    ok = lambda r: isinstance(r, dict) and not r.get("error") and not r.get("status")
    if ok(a) and ok(b):
        sa, sb = check_score(seg_result["checks"]["model_a"]), check_score(seg_result["checks"]["model_b"])
        if sa != sb:
            return ("B", b) if sb > sa else ("A", a)
        return ("B", b) if completeness(b) > completeness(a) * 1.1 else ("A", a)  # tie: the fuller table wins
    return ("A", a) if ok(a) else (("B", b) if ok(b) else (None, None))


def normalize_result(r):
    """Models sometimes return table headers as plain strings: give them roles from the data so every check can run."""
    if not isinstance(r, dict) or r.get("error"):
        return r
    if isinstance(r.get("tables"), list):  # "tables" without columns are stray text, not tables
        r["tables"] = [t for t in r["tables"] if isinstance(t, dict) and isinstance(t.get("columns"), list) and t["columns"]]
    for p in r.get("parties") or []:  # IDs printed in brackets, e.g. "(1340607-U)": keep the ID itself
        for k in ("registration_no", "tax_id"):
            if isinstance(p, dict) and isinstance(p.get(k), str):
                p[k] = re.sub(r"^\((.*)\)$", r"\1", p[k].strip()).strip()
    for t in r.get("tables") or []:
        if not isinstance(t.get("columns"), list) or not any(isinstance(c, str) for c in t["columns"]):
            continue
        rows = [x for x in t.get("rows") or [] if isinstance(x, list)]
        cell = lambda row, i: str(row[i] if i < len(row) and row[i] is not None else "").strip()
        cols = []
        for i, c in enumerate(t["columns"]):
            if isinstance(c, dict):
                cols.append(c); continue
            vals = [v for v in (cell(row, i) for row in rows) if v]
            if not vals: role = "text"
            elif len(vals) == len(rows) and all(v == str(k + 1) for k, v in enumerate(vals)): role = "id"
            elif all(num(v) is not None and re.fullmatch(r"[\s(\-]*[A-Z$€£¥]{0,3}\s*[\d,.]+\s*\)?\s*(CR|DR|-)?", v, re.I) for v in vals):
                role = "id" if all(re.fullmatch(r"\d{6,}", v) for v in vals) else "number"
            elif all(norm_date(v) for v in vals): role = "date"
            elif re.search(r"\d", "".join(vals)) and all(re.fullmatch(r"[\w\-/.#]+", v) for v in vals): role = "id"
            else: role = "text"
            cols.append({"name": str(c if c is not None else ""), "role": role})
        nums = [i for i, c in enumerate(cols) if c["role"] == "number"]
        for i in nums:  # a numeric column equal to the sum of the others on every row is the row total
            others = [j for j in nums if j != i]
            if others and rows and all(num(cell(row, i)) is None or abs(num(cell(row, i)) - sum(num(cell(row, j)) or 0 for j in others)) < 0.011 for row in rows):
                cols[i]["role"] = "row_total"; break
        t["columns"] = cols
    return r


# ----------------------------------------------------------------- pipeline
@dataclass
class Options:
    a: ModelConfig = field(default_factory=ModelConfig)
    b: ModelConfig = field(default_factory=lambda: ModelConfig(type="anthropic", base="", model="claude-sonnet-5-5"))
    enhance: bool = True
    tiles: bool = False
    concurrency: int = 2
    pdf_text: str = "auto"          # auto | image


def _local_table_result(seg: Segment):
    rows = seg.raw_rows
    head = [str(h) for h in rows[0]]
    is_num = lambda v: v != "" and num(v) is not None and re.fullmatch(r"[\d.,\s-]+", str(v)) is not None
    cols = [{"name": h or f"Column {i+1}", "role": "number" if all(r[i] == "" or is_num(r[i]) for r in rows[1:] if i < len(r)) and any(i < len(r) and is_num(r[i]) for r in rows[1:]) else "text"} for i, h in enumerate(head)]
    body, total = rows[1:], None
    if body and re.search(r"total|jumlah|合计|总计", str(body[-1][0]), re.I):
        total = body.pop()
    return {"model": "local", "document_type": "other", "title": seg.source, "tables": [{"name": seg.source, "columns": cols, "rows": body, "total_row": total}], "fields": [], "parties": []}


async def process_segment(client, seg: Segment, o: Options):
    hint = ""
    if seg.kind == "text":
        hint = (f"\n\nThere is NO image. The document content below was extracted from a {seg.source} (cells separated by \" | \", one row per line"
                f"{'; content truncated' if seg.truncated else ''}). Apply the same rules to this text; rotation_degrees = \"0\".\n<<<DOCUMENT\n{norm_text(seg.text)}\nDOCUMENT>>>")
    elif seg.text:
        hint = f"\n\nPDF text layer for reference:\n{norm_text(seg.text)[:6000]}"
    imgs = prep_images(seg.b64, o.enhance, o.tiles) if seg.b64 and not seg.text_only else None
    text = PROMPT + hint + (MULTI_NOTE.format(n=len(imgs)) if imgs and len(imgs) > 1 else "")
    local = local_codes(seg)
    if not o.a.enabled and not o.b.enabled and seg.raw_rows:
        ra, rb = _local_table_result(seg), {"status": "skipped"}
    else:
        run = lambda c: safe(call_model(client, c, text, imgs)) if c.enabled else asyncio.sleep(0, {"status": "skipped"})
        same = o.a.enabled and o.b.enabled and o.a.type == o.b.type and clean_base(o.a.base) == clean_base(o.b.base)
        if same:  # same relay: call B after A so its concurrency limit doesn't reject B
            ra = await run(o.a)
            rb = await run(o.b)
        else:
            ra, rb = await asyncio.gather(run(o.a), run(o.b))
        # the model says the page is turned: straighten it and read once more (sideways tables shift columns/rows)
        rd = next((d for d in (_int((r or {}).get("rotation_degrees")) for r in (ra, rb) if isinstance(r, dict)) if d in (90, 180, 270)), None)
        if rd and seg.b64 and not seg.text_only and not seg.straightened:
            seg.b64, seg.straightened = jpeg_b64(b64_image(seg.b64).rotate(-rd, expand=True)), rd
            first = {"model_a": ra.get("__usage") if isinstance(ra, dict) else None, "model_b": rb.get("__usage") if isinstance(rb, dict) else None}
            res = await process_segment(client, seg, o)
            for k, u in first.items():
                if u:
                    cur = res["usage"].get(k) or {"in": 0, "out": 0}
                    res["usage"][k] = {"in": (cur.get("in") or 0) + (u.get("in") or 0), "out": (cur.get("out") or 0) + (u.get("out") or 0)}
            res["straightened"] = rd
            return res
    ra = normalize_result(ra if isinstance(ra, dict) else {"error": "invalid model output"})
    rb = normalize_result(rb if isinstance(rb, dict) else {"error": "invalid model output"})
    align_by_label(ra, rb)
    usage = {"model_a": ra.pop("__usage", None), "model_b": rb.pop("__usage", None)}
    res = {"segment_identifier": seg.id, "source_kind": seg.kind, "usage": usage,
           "image_size": list(b64_image(seg.b64).size) if seg.b64 and not seg.text_only else None, "local_extraction": local,
           "provider_responses": {"model_a": {"model": o.a.model, **ra} if "model" not in ra else ra, "model_b": {"model": o.b.model, **rb}},
           "checks": {"model_a": verify(ra), "model_b": verify(rb)}, "cross_verification": compare(ra, rb)}
    who, final = pick_model(res)
    res["final_source"], res["final_result"], res["manual_edits"], res["accepted"] = who, final, [], []
    res["merged"] = merge_from(final, who, (res.get("cross_verification") or {}).get("discrepancies")) if final is not None else []
    if seg.text and seg.b64:  # PDF page with a real text layer: a third, deterministic check of the model's values
        res["grounding"] = grounding(norm_text(seg.text), final)
    return res


async def process_file(client, path: Path, o: Options, rotate=0):
    t0 = time.time()
    segs = ingest(path, rotate, o.pdf_text)
    out = []
    for i in range(0, len(segs), 3):
        out += await asyncio.gather(*(process_segment(client, s, o) for s in segs[i:i + 3]))
    return {"source_document": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "processed_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "elapsed_ms": int((time.time() - t0) * 1000),
            "total_extracted_segments": len(out), "segments": out}


def status_of(result) -> str:
    """ok | warn | bad | err — same idea as the queue badges in the web app."""
    if result.get("error"):
        return "err"
    bad = warn = 0
    for sg in result["segments"]:
        pr = sg["provider_responses"]
        if all(pr[k].get("error") or pr[k].get("status") for k in ("model_a", "model_b")) and pr["model_a"].get("model") != "local":
            return "err"
        for k in ("model_a", "model_b"):
            c = sg["checks"][k]
            if c:
                bad += sum(not x["ok"] for x in c["checks"])
            warn += len(pr[k].get("low_confidence_fields") or [])
        if sg["cross_verification"]:
            warn += len(sg["cross_verification"]["discrepancies"])
    return "bad" if bad else "warn" if warn else "ok"


async def run_batch(paths, o: Options, transport=None, on_done=None):
    seen = {}
    sem = asyncio.Semaphore(max(1, o.concurrency))
    async with httpx.AsyncClient(transport=transport) as client:
        async def one(p):
            h = hashlib.sha256(p.read_bytes()).hexdigest()
            if h in seen:
                return {"source_document": p.name, "skipped": f"identical to {seen[h]}"}
            seen[h] = p.name
            async with sem:
                try:
                    r = await process_file(client, p, o)
                except Exception as e:
                    r = {"source_document": p.name, "error": str(e), "segments": []}
            r["status"] = status_of(r)
            if on_done:
                on_done(r)
            return r
        async def skip(p):
            return {"source_document": p.name, "skipped": "unsupported file type"}
        # results keep the input order; identical files are detected in input order too
        results = await asyncio.gather(*(skip(p) if kind_of(p) is None else one(p) for p in paths))
    return list(results)


# ----------------------------------------------------------------- reconciliation & duplicates (same rules as the web app)
INV_COL = re.compile(r"(invoice|inv|bill)\s*\.?\s*(no|number|num|#)|no\.?\s*(invois|inv)|发票号|发票编号|單號|单号", re.I)
DATE_COL = re.compile(r"date|tarikh|日期", re.I)
NAME_COL = re.compile(r"\bname\b|customer|nama|pelanggan|名称|客户", re.I)
SHIP_COL = re.compile(r"ship.?to|deliver.?to", re.I)
norm_no = lambda v: re.sub(r"[^A-Z0-9]", "", str(v if v is not None else "").upper())


def ocr_canon(k: str) -> str:
    for a, b in (("[OQD]", "0"), ("[IL|]", "1"), ("S", "5"), ("B", "8"), ("Z", "2"), ("G", "6")):
        k = re.sub(a, b, k)
    return k


def norm_date(v):
    m = re.search(r"(\d{1,4})[/.\-](\d{1,2})[/.\-](\d{1,4})", str(v or ""))
    if not m:
        return None
    a, b, c = m.groups()
    y, mo, d = (a, b, c) if len(a) == 4 else (("20" + c) if len(c) == 2 else c, b, a)
    return f"{y}-{int(mo):02d}-{int(d):02d}"


def reconcile(results):
    registers, invoices = [], []
    for res in results:
        docs = [sg.get("final_result") for sg in res.get("segments", []) if sg.get("final_result")]
        reg_rows = []
        for r in docs:
            for tb in r.get("tables") or []:
                cols = tb.get("columns") or []
                ci = next((i for i, c in enumerate(cols) if INV_COL.search((c or {}).get("name") or "")), -1)
                if ci < 0:
                    continue
                f = lambda rx, extra=lambda c: True: next((i for i, c in enumerate(cols) if rx.search((c or {}).get("name") or "") and extra(c)), -1)
                di, ni = f(DATE_COL), f(NAME_COL, lambda c: not INV_COL.search(c.get("name") or ""))
                ni = ni if ni >= 0 else f(SHIP_COL)
                ai = next((i for i, c in enumerate(cols) if (c or {}).get("role") == "amount"), -1)
                for row in tb.get("rows") or []:
                    if isinstance(row, list) and ci < len(row) and norm_no(row[ci]):
                        g = lambda i: row[i] if 0 <= i < len(row) else None
                        reg_rows.append({"no": row[ci], "key": norm_no(row[ci]), "date": g(di), "name": g(ni), "amount": g(ai)})
        if (any(r.get("document_type") == "invoice_register" for r in docs) or len(reg_rows) >= 2) and reg_rows:
            registers.append({"file": res["source_document"], "no": docs[0].get("document_number") if docs else None, "rows": reg_rows})
        elif docs and norm_no(docs[0].get("document_number")):
            r = docs[0]
            invoices.append({"file": res["source_document"], "no": r.get("document_number"), "key": norm_no(r.get("document_number")),
                             "date": r.get("document_date"), "total": r.get("grand_total")})
    by_key = {}
    for x in invoices:
        by_key.setdefault(x["key"], []).append(x)
    dups = [{"no": v[0]["no"], "files": [x["file"] for x in v]} for v in by_key.values() if len(v) > 1]
    used = set()
    rows = [r for reg in registers for r in reg["rows"]]
    for row in rows:
        m = next((x for x in invoices if x["key"] == row["key"]), None)
        row["match"], row["kind"] = m, "match" if m else "missing"
        if m: used.add(m["file"])
    for row in rows:
        if row["match"]:
            continue
        c = ocr_canon(row["key"])
        m = next((x for x in invoices if x["file"] not in used and len(x["key"]) >= 4 and ocr_canon(x["key"]) == c), None)
        if m:
            row["match"], row["kind"] = m, "near"; used.add(m["file"])
    for row in rows:
        m = row["match"]
        row["date_bad"] = bool(m and norm_date(row["date"]) and norm_date(m["date"]) and norm_date(row["date"]) != norm_date(m["date"]))
        row["amount_bad"] = bool(m and num(row["amount"]) is not None and num(m["total"]) is not None and abs(num(row["amount"]) - num(m["total"])) > 0.011)
        row["status"] = ("missing" if row["kind"] == "missing" else "possible_match" if row["kind"] == "near"
                         else "date_mismatch" if row["date_bad"] else "amount_mismatch" if row["amount_bad"] else "ok")
    for reg in registers:  # a register with no uploaded invoice of its own is informational, not a list of errors
        reg["related"] = any(r["match"] for r in reg["rows"])
        if not reg["related"]:
            for r in reg["rows"]:
                r["status"] = "not_uploaded"
    extra = [x for x in invoices if x["file"] not in used] if any(g["related"] for g in registers) else []
    return {"registers": registers, "extra": extra, "duplicates": dups}


# ----------------------------------------------------------------- outputs / CLI
def flat_rows(file, page, r):
    """Every extracted value as rows: file, page, section, item, label, value."""
    out = []
    def put(sec, item, label, v):
        if v not in (None, ""):
            out.append([file, page, sec, item, label, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
    for k in ("document_type", "title", "document_number", "document_date", "currency", "grand_total", "amount_in_words"):
        put("document", "", k, r.get(k))
    for i, p in enumerate(x for x in r.get("parties") or [] if isinstance(x, dict)):
        for k in ("name", "registration_no", "tax_id", "address", "contact"):
            put("party", f"{p.get('role') or 'party'} {i + 1}", k, p.get(k))
    for i, x in enumerate(r.get("fields") or []):
        put("field", i + 1, (x or {}).get("label"), (x or {}).get("value"))
    for i, x in enumerate(r.get("totals") or []):
        put("total", i + 1, (x or {}).get("label"), (x or {}).get("value"))
    for ti, t in enumerate(r.get("tables") or []):
        cols, tn = [c.get("name") for c in t.get("columns") or []], t.get("name") or f"Table {ti + 1}"
        col = lambda ci: cols[ci] if ci < len(cols) and cols[ci] else f"col {ci + 1}"
        for ri, row in enumerate(t.get("rows") or []):
            for ci, v in enumerate(row or []):
                put(f"table: {tn}", ri + 1, col(ci), v)
        if isinstance(t.get("total_row"), list):
            for ci, v in enumerate(t["total_row"]):
                put(f"table: {tn}", "TOTAL", col(ci), v)
    for sec, key in (("stamp", "stamps_and_chops"), ("handwriting", "handwritten_notes")):
        for i, x in enumerate(r.get(key) or []):
            put(sec, i + 1, (x.get("position") or "") if isinstance(x, dict) else "", x.get("text") if isinstance(x, dict) else x)
    return out


def write_outputs(results, out: Path):
    out.mkdir(parents=True, exist_ok=True)
    rc = reconcile([r for r in results if r.get("segments")])
    (out / "results.json").write_text(json.dumps({"exported_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "documents": results, "reconciliation": rc}, ensure_ascii=False, indent=2), encoding="utf-8")
    dup_of = {f: ", ".join(x for x in d["files"] if x != f) for d in rc["duplicates"] for f in d["files"]}
    with open(out / "fields.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["file", "page", "section", "item", "label", "value"])
        for r in results:
            for sg in r.get("segments", []):
                if sg.get("final_result"):
                    w.writerows(flat_rows(r["source_document"], sg["segment_identifier"], sg["final_result"]))
    with open(out / "summary.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["file", "status", "duplicate_of", "tokens_in", "tokens_out", "document_type", "document_number", "document_date", "currency", "grand_total", "table_rows", "note"])
        for r in results:
            fr = next((sg["final_result"] for sg in r.get("segments", []) if sg.get("final_result")), {}) or {}
            us = [u for sg in r.get("segments", []) for u in (sg.get("usage") or {}).values() if u]
            w.writerow([r["source_document"], r.get("status", "skipped"), dup_of.get(r["source_document"], ""), sum(u["in"] for u in us), sum(u["out"] for u in us), fr.get("document_type"), fr.get("document_number"), fr.get("document_date"),
                        fr.get("currency"), fr.get("grand_total"), sum(len(t.get("rows") or []) for t in fr.get("tables") or []), r.get("error") or r.get("skipped") or ""])
    with open(out / "tables.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for r in results:
            for sg in r.get("segments", []):
                fr = sg.get("final_result") or {}
                for ti, t in enumerate(fr.get("tables") or []):
                    pre = lambda k: [r["source_document"], sg["segment_identifier"], fr.get("document_number"), t.get("name") or f"Table {ti+1}", k]
                    w.writerow(["file", "page", "document_number", "table", "row"] + [c.get("name") for c in t.get("columns") or []])
                    for ri, row in enumerate(t.get("rows") or []):
                        w.writerow(pre(ri + 1) + list(row))
                    if isinstance(t.get("total_row"), list):
                        w.writerow(pre("TOTAL") + list(t["total_row"]))
                    w.writerow([])
    with open(out / "reconcile.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["register_file", "invoice_no", "register_date", "name", "matched_file", "matched_invoice_no", "invoice_date", "invoice_total", "status"])
        for reg in rc["registers"]:
            for row in reg["rows"]:
                m = row["match"] or {}
                w.writerow([reg["file"], row["no"], row["date"], row["name"], m.get("file"), m.get("no"), m.get("date"), m.get("total"), row["status"]])
        for x in rc["extra"]:
            w.writerow(["", x["no"], "", "", x["file"], x["no"], x["date"], x["total"], "not_in_register"])
        for d in rc["duplicates"]:
            w.writerow(["", d["no"], "", "", " | ".join(d["files"]), d["no"], "", "", "duplicate"])
    return rc


def main(argv=None, transport=None):
    ap = argparse.ArgumentParser(description="InvoiceLens command-line extractor")
    ap.add_argument("files", nargs="+", type=Path)
    for s, d_type, d_base, d_model in (("a", "openai", "https://api.openai.com/v1", "gpt-4o"), ("b", "anthropic", "", "claude-sonnet-5-5")):
        ap.add_argument(f"--{s}-type", default=os.getenv(f"IL_{s.upper()}_TYPE", d_type), choices=["openai", "anthropic"])
        ap.add_argument(f"--{s}-base", default=os.getenv(f"IL_{s.upper()}_BASE", d_base))
        ap.add_argument(f"--{s}-model", default=os.getenv(f"IL_{s.upper()}_MODEL", d_model))
        ap.add_argument(f"--{s}-key", default=os.getenv(f"IL_{s.upper()}_KEY", ""), help=f"or set IL_{s.upper()}_KEY")
        ap.add_argument(f"--{s}-no-json", action="store_true", help="disable JSON mode for this OpenAI-compatible provider")
    ap.add_argument("--no-enhance", action="store_true", help="disable small-text enhancement")
    ap.add_argument("--tiles", action="store_true", help="multi-scale reading (full page + two enlarged halves)")
    ap.add_argument("--concurrency", type=int, default=2)
    ap.add_argument("--pdf-image", action="store_true", help="always send PDF pages as images (default: use the text layer when present)")
    ap.add_argument("--timeout", type=float, default=MODEL_TIMEOUT)
    ap.add_argument("--max-tokens", type=int, default=int(os.getenv("IL_MAX_TOKENS", "32000")), help="max output tokens per model call (thinking models need a high value)")
    ap.add_argument("--out", type=Path, default=Path("invoicelens-out"))
    a = ap.parse_args(argv)
    mk = lambda s: ModelConfig(getattr(a, f"{s}_type"), getattr(a, f"{s}_base"), getattr(a, f"{s}_model"), getattr(a, f"{s}_key"), not getattr(a, f"{s}_no_json"), a.timeout, a.max_tokens)
    o = Options(mk("a"), mk("b"), not a.no_enhance, a.tiles, a.concurrency, "image" if a.pdf_image else "auto")
    if not o.a.enabled and not o.b.enabled:
        print("No model key configured — only local parsing (CSV/Excel tables, barcodes).", file=sys.stderr)
    files = [p for f in a.files for p in (sorted(x for x in f.iterdir() if x.is_file()) if f.is_dir() else [f])]
    results = asyncio.run(run_batch(files, o, transport, on_done=lambda r: print(f"{r['status']:5} {r['source_document']}", file=sys.stderr)))
    rc = write_outputs(results, a.out)
    bad = [r for r in results if r.get("status") in ("bad", "err")]
    print(f"{len(results)} file(s) → {a.out}/ (results.json, summary.csv, tables.csv, reconcile.csv); "
          f"{len(bad)} need review; {sum(r['status'] not in ('ok', 'not_uploaded') for reg in rc['registers'] for r in reg['rows'])} reconciliation issue(s)", file=sys.stderr)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
