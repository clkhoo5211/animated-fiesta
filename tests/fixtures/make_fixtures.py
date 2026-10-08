"""Regenerate synthetic test fixtures (no real customer data).
Requires: pip install pillow qrcode pillow-heif openpyxl python-docx
Run: python tests/fixtures/make_fixtures.py
"""
import os
from PIL import Image, ImageDraw
import qrcode

D = os.path.dirname(os.path.abspath(__file__))
p = lambda n: os.path.join(D, n)

# --- synthetic invoice image with a QR code ---
im = Image.new("RGB", (800, 1100), "white")
d = ImageDraw.Draw(im)
lines = ["DEMO TRADING SDN BHD (000000-X)", "No. 1 Jalan Contoh, 50000 Kuala Lumpur", "",
         "INVOICE  No: INV-TEST-0001   Date: 01/10/2026", "Bill to: SAMPLE STORE SDN BHD", "",
         "Item            Qty   Unit Price   Amount", "Widget A          6       20.86     125.16",
         "Widget B         10        4.96      49.60", "Service           1       50.00      50.00", "",
         "Total Payable (RM)                  224.76"]
for i, t in enumerate(lines):
    d.text((40, 40 + i * 28), t, fill="black")
qr = qrcode.make("INV-TEST-0001|TOTAL=224.76").convert("RGB").resize((260, 260))
im.paste(qr, (480, 780))
im.save(p("invoice.jpg"), quality=90)
im.save(p("invoice.pdf"))
im.rotate(90, expand=True).save(p("rotated.jpg"), quality=90)

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
    im.save(p("photo.heic"), quality=70)
except ImportError:
    print("pillow-heif not installed: photo.heic not regenerated")

# --- spreadsheet / csv / word ---
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Invoice"
ws.append(["Invoice No", "INV-7788"]); ws.append(["Date", "2026-10-01"]); ws.append([])
ws.append(["Item", "Qty", "Unit Price", "Amount"])
for r in [["Widget A", 2, 10.5, 21.0], ["Widget B", 3, 4.0, 12.0], ["Service", 1, 50, 50]]:
    ws.append(r)
ws.append(["Total", None, None, 83.0])
wb.create_sheet("Notes").append(["Remark", "Deliver before Friday"])
wb.save(p("sheet.xlsx"))

open(p("data.csv"), "w").write("Delivery No,Name,CTN\n1000000001,SHOP ALPHA,1\n1000000002,SHOP BETA,28\nTotal,,29\n")

import docx
doc = docx.Document(); doc.add_heading("Quotation Q-2026-15", 1)
doc.add_paragraph("Customer: ACME Sdn Bhd   Date: 5 Oct 2026")
t = doc.add_table(rows=1, cols=3)
for i, h in enumerate(["Item", "Qty", "Amount"]):
    t.rows[0].cells[i].text = h
for row in [("Design", "1", "1500.00"), ("Hosting", "12", "600.00")]:
    c = t.add_row().cells
    for i, v in enumerate(row):
        c[i].text = v
doc.add_paragraph("Total: 2100.00"); doc.save(p("quote.docx"))
open(p("register.csv"), "w").write("Delivery No,Name,Invoice Date,Invoice No.,CTN-1,Total\n1000000001,SHOP ALPHA,01/10/2026,INV-1001,1,1\n1000000002,SHOP BETA,01/10/2026,INV-1002,2,2\n1000000003,SHOP GAMMA,02/10/2026,INV-1003,3,3\n1000000004,SHOP DELTA,02/10/2026,INV-1004,4,4\nTotal,,,,10,10\n")
open(p("note.txt"), "w").write("unsupported file\n")
print("fixtures written to", D)
